import os
import asyncio
from fastapi import HTTPException
from langchain_community.document_loaders import TextLoader, UnstructuredPowerPointLoader
from langchain.document_loaders import PyMuPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain.vectorstores import Chroma
import asyncpg
from dotenv import load_dotenv
import nltk
from pptx import Presentation
from langchain.schema import Document
import subprocess
from config import DATABASE_CONFIG
nltk.download('punkt')

load_dotenv()
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")


async def get_file_name_from_db(file_id: str) -> str:
    query = """
        SELECT file_nm
        FROM FILE_MST
        WHERE file_id = $1
    """
    
    try:
        print(f"Fetching file name for file_id: {file_id}")  # 디버깅 로그
        conn = await asyncpg.connect(**DATABASE_CONFIG)
        try:
            result = await conn.fetchval(query, file_id)
            print(f"Query result for file_id {file_id}: {result}")  # 디버깅 로그
            return result
        finally:
            await conn.close()
    except Exception as e:
        print(f"Error fetching file name from DB: {e}")  # 디버깅 로그
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")

def preprocess_text_with_category(text, file_type_str, title):
    """
    파일 내용 앞에 종류 정보를 추가하여 임베딩할 때 활용할 수 있도록 변환
    """
    category_text = f"이 문서는 교수자가 [{file_type_str}] 카테고리에 업로드한 내용입니다. 참고한 파일의 제목은 [{title}]입니다. \n\n"
    return category_text + text

def load_text_file(file_path, title=None, file_type_str=None):
    """
    텍스트 파일 로드 및 처리
    """
    loader = TextLoader(file_path)
    documents = loader.load()

    # 메타데이터 추가 + 종류 정보 포함한 텍스트 변환
    for doc in documents:
        doc.page_content = preprocess_text_with_category(doc.page_content, file_type_str, title)
        if title:
            doc.metadata["title"] = title
        if file_type_str:
            doc.metadata["종류"] = file_type_str
    return documents

def load_pdf_file(file_path, title=None, file_type_str=None):
    """
    PDF 파일 로드 및 처리
    """
    loader = PyMuPDFLoader(file_path)
    documents = loader.load()

    for doc in documents:
        doc.page_content = preprocess_text_with_category(doc.page_content, file_type_str, title)
        if title:
            doc.metadata["title"] = title
        if file_type_str:
            doc.metadata["종류"] = file_type_str
    return documents

def load_ppt_file(file_path, title=None, file_type_str=None):
    """
    파워포인트 파일 로드 및 텍스트 추출
    """
    def extract_text_from_pptx(file_path):
        presentation = Presentation(file_path)
        texts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            texts.append(text)
        return texts
    
    # 파일 확장자가 .ppt일 경우 변환 후 사용 ##0214
    if file_path.endswith(".ppt"):
        file_path = convert_ppt_to_pptx(file_path)

    texts = extract_text_from_pptx(file_path)
    documents = []
    for text in texts:
        text = preprocess_text_with_category(text, file_type_str, title)  # 종류 정보 추가
        metadata = {
            "source": file_path,
            "file_type": "pptx"
        }
        if title:
            metadata["title"] = title
        if file_type_str:
            metadata["종류"] = file_type_str

        document = Document(
            page_content=text,
            metadata=metadata
        )
        documents.append(document)

    return documents
def convert_ppt_to_pptx(ppt_path):
    """LibreOffice를 사용하여 PPT를 PPTX로 변환"""
    output_path = ppt_path.replace(".ppt", ".pptx")
    
    command = [
        "libreoffice", "--headless", "--convert-to", "pptx", ppt_path, "--outdir", os.path.dirname(ppt_path)
    ]
    subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

    if os.path.exists(output_path):
        return output_path  # 변환 성공
    else:
        raise FileNotFoundError(f"변환 실패: {ppt_path}")
    
async def load_file(file_path, file_type_cd, file_id):
    """
    파일을 로드하고 태그를 추가하는 함수
    """
    type_mapping = {
        "S": "강의일정",
        "N": "공지사항",
        "M": "수업자료",
        "V": "강의영상"
    }
    file_type_str = type_mapping.get(file_type_cd, "기타")
    
    # DB에서 파일명 가져오기
    title = await get_file_name_from_db(file_id)
    if not title:
        print(f"Database did not return a title for file_id {file_id}. Using file name from path.")
        title = os.path.basename(file_path)
    
    print(f"Loading file with title: {title} and type: {file_type_str}")  # 디버깅 로그 추가

    ext = os.path.splitext(file_path)[-1].lower()
    
    if ext == ".txt":
        documents = load_text_file(file_path, title, file_type_str)
    elif ext in [".pptx", ".ppt"]: 
        documents = load_ppt_file(file_path, title, file_type_str)
    elif ext == ".pdf":
        documents = load_pdf_file(file_path, title, file_type_str)
    else:
        raise ValueError(f"지원되지 않는 파일 형식: {ext}")

    return documents

async def update_status_in_db(status, dt_tag, file_id):
    """
    데이터베이스에서 파일 상태 업데이트
    """
    query = f"""
        UPDATE FILE_MST
        SET UPLOAD_STATUS = $1,
        {dt_tag} = CURRENT_TIMESTAMP AT TIME ZONE 'Asia/Seoul'
        WHERE FILE_ID = $2
    """

    try:
        conn = await asyncpg.connect(**DATABASE_CONFIG)
        try:
            await conn.execute(query, status, file_id)
        finally:
            await conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")

    print(f"DB 업데이트: {file_id} - {status}")

async def process_and_embed_files(cls_id: str, files_info: list, notify_filestatus_toclients):
    """
    파일 처리 및 임베딩 수행
    """
    base_path = f"./uploaded_files/{cls_id}"
    chroma_db_path = os.path.join("./db/chromadb", cls_id)

    # 경로가 없으면 디렉토리 생성
    os.makedirs(base_path, exist_ok=True)
    os.makedirs(chroma_db_path, exist_ok=True)
    
    # 디렉토리 권한 설정
    os.chmod(base_path, 0o775)
    os.chmod(chroma_db_path, 0o775)
    
    # ChromaDB 관련 파일들의 권한도 설정
    for root, dirs, files in os.walk(chroma_db_path):
        for dir in dirs:
            os.chmod(os.path.join(root, dir), 0o775)
        for file in files:
            os.chmod(os.path.join(root, file), 0o664)
    
    for file_info in files_info:
        file_name = file_info["file_name"]
        file_id = file_info["file_id"]
        file_type_cd = file_info["file_type_cd"]
        file_path = os.path.join(base_path, file_name)

        try:
            # 파일 로드 및 태그 추가
            documents = await load_file(file_path, file_type_cd, file_id)

            # 로드 성공 처리
            await update_status_in_db("EP02", 'EMB_START_DT', file_id)
            await notify_filestatus_toclients("파일 상태 업데이트")
            print(f"{file_id} 파일 로드 성공!")

            # 텍스트 분할
            text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
            docs = text_splitter.split_documents(documents)

            for doc in docs:
                doc.metadata["title"] = documents[0].metadata.get("title", "Unknown")
                doc.metadata["종류"] = documents[0].metadata.get("종류", "Unknown")

            # 임베딩 생성 및 저장
            embeddings = OpenAIEmbeddings(model="text-embedding-3-large")
            vector_store = Chroma.from_documents(
                docs,
                embeddings,
                persist_directory=chroma_db_path,
                collection_name=cls_id
            )
            vector_store.persist()
            
            # 새로 생성된 ChromaDB 파일들의 권한 설정
            for root, dirs, files in os.walk(chroma_db_path):
                for dir in dirs:
                    os.chmod(os.path.join(root, dir), 0o775)
                for file in files:
                    os.chmod(os.path.join(root, file), 0o664)

            # 임베딩 성공 처리
            await update_status_in_db("EP03", 'EMB_COMP_DT', file_id)
            await notify_filestatus_toclients("파일 상태 업데이트")
            print(f"{file_id} 파일 임베딩 성공!")

        except Exception as e:
            # 임베딩 실패 처리
            await update_status_in_db("EP04", 'EMB_FAIL_DT', file_id)
            await notify_filestatus_toclients("파일 상태 업데이트")
            print(f"{file_id} 임베딩 실패: {str(e)}")

# 예제 실행
if __name__ == "__main__":
    json_data = {
        "cls_id": "2025-1-506819-01",
        "files_info": [
            {
                "file_name": "2025-1-506819-01-8a6728.pdf",
                "file_id": "2025-1-506819-01-8a6728",
                "file_type_cd": "S"
            }
            # {
            #     "file_name": "FILE_0000000019.pdf",
            #     "file_id": "FILE_0000000019",
            #     "file_type_cd": "N"
            # }
        ]
    }

    asyncio.run(process_and_embed_files(json_data["cls_id"], json_data["files_info"]))